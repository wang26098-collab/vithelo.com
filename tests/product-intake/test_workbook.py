import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[2]
MODULE_DIR = ROOT / "scripts" / "product-intake"
sys.path.insert(0, str(MODULE_DIR))

from constants import (  # noqa: E402
    CLAIM_COLUMNS,
    IMAGE_COLUMNS,
    PARAMETER_COLUMNS,
    PRODUCT_COLUMNS,
    SHEET_NAMES,
)
from workbook import WorkbookContractError, read_workbook  # noqa: E402


def write_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[object]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)


def make_workbook(path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.create_sheet("填写说明")
    write_sheet(
        workbook,
        "产品主表",
        PRODUCT_COLUMNS,
        [
            [
                "示例不导入",
                "EXAMPLE-001",
                "Gummies",
                "Example Gummies Product for Private Label Nutrition",
                "示例软糖",
                "This row explains how to complete the template and must never be imported into the product catalog.",
                "Dietary supplement",
                "Mixed berry",
                "Example formula direction",
                "60 g",
                "30 gummies",
                "Pouch",
                "Berry red",
                "",
                "",
                "",
                "",
                "待确认",
                "示例",
                "产品主表 row 2",
                "示例，不导入",
            ],
            [
                "正式资料",
                "VIT-GUM-001",
                "Gummies",
                "Mixed Berry Gummies for Private Label Nutrition Projects",
                "混合莓软糖",
                "Source-provided private-label product direction for structured review; formula, packaging and production specifications remain pending verification.",
                "Dietary supplement",
                "Mixed berry",
                "",
                "60 g",
                "30 gummies",
                "Pouch",
                "Berry red",
                "",
                "",
                "Flavor; shape; packaging",
                "",
                "可公开",
                "products.xlsx",
                "产品主表 row 3",
                "",
            ],
            [
                "正式资料",
                "VIT-CAP-001",
                "Hard Capsules",
                "Hard Capsules for Private Label Nutrition Projects",
                "硬胶囊",
                "Source-provided capsule product direction for structured review; formula, packaging and production specifications remain pending verification.",
                "Dietary supplement",
                "",
                "",
                "",
                "60 capsules",
                "Bottle",
                "",
                "",
                "",
                "Formula; packaging",
                "",
                "可公开",
                "products.xlsx",
                "产品主表 row 4",
                "",
            ],
        ],
    )
    write_sheet(
        workbook,
        "产品参数",
        PARAMETER_COLUMNS,
        [["VIT-GUM-001", "Flavor", "Mixed berry", "口味", "混合莓", 1, "可公开", "products.xlsx", "产品参数 row 2"]],
    )
    write_sheet(
        workbook,
        "图片清单",
        IMAGE_COLUMNS,
        [["VIT-GUM-001", "main.png", "主图", 1, "VITHELO mixed berry gummies package", "User supplied", "可公开", ""]],
    )
    write_sheet(
        workbook,
        "宣称与证据",
        CLAIM_COLUMNS,
        [["VIT-GUM-001", "Mixed berry flavor", "混合莓口味", "产品描述", "products.xlsx", "产品主表 row 3", "VIT-GUM-001", "Source-provided specification", "可公开", "可采用"]],
    )
    workbook.save(path)


class WorkbookContractTest(unittest.TestCase):
    def test_reads_multiple_products_and_ignores_example_rows(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "products.xlsx"
            make_workbook(path)

            result = read_workbook(path)

            self.assertEqual(result.sheet_names, SHEET_NAMES)
            self.assertEqual([item["产品编号"] for item in result.products], ["VIT-GUM-001", "VIT-CAP-001"])
            self.assertEqual(result.parameters[0]["产品编号"], "VIT-GUM-001")
            self.assertEqual(result.images[0]["图片用途"], "主图")
            self.assertEqual(result.claims[0]["审核结论"], "可采用")

    def test_rejects_a_changed_header(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "products.xlsx"
            make_workbook(path)
            from openpyxl import load_workbook

            workbook = load_workbook(path)
            workbook["产品主表"]["A1"] = "错误表头"
            workbook.save(path)

            with self.assertRaisesRegex(WorkbookContractError, "产品主表"):
                read_workbook(path)


if __name__ == "__main__":
    unittest.main()
