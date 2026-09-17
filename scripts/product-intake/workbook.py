from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from constants import SHEET_NAMES, TABLE_COLUMNS


class WorkbookContractError(ValueError):
    pass


@dataclass(frozen=True)
class WorkbookData:
    source: Path
    sheet_names: list[str]
    products: list[dict[str, Any]]
    parameters: list[dict[str, Any]]
    images: list[dict[str, Any]]
    claims: list[dict[str, Any]]


def _clean(value: Any) -> Any:
    return value.strip() if isinstance(value, str) else value


def _read_rows(sheet, expected_headers: list[str]) -> list[dict[str, Any]]:
    actual_headers = [_clean(cell.value) for cell in sheet[1]][: len(expected_headers)]
    if actual_headers != expected_headers:
        raise WorkbookContractError(
            f"{sheet.title} 表头不匹配。预期 {expected_headers}，实际 {actual_headers}。"
        )

    rows: list[dict[str, Any]] = []
    for values in sheet.iter_rows(min_row=2, max_col=len(expected_headers), values_only=True):
        cleaned = [_clean(value) for value in values]
        if not any(value not in (None, "") for value in cleaned):
            continue
        row = dict(zip(expected_headers, cleaned, strict=True))
        product_id = row.get("产品编号")
        if isinstance(product_id, str) and product_id.startswith("EXAMPLE-"):
            continue
        rows.append(row)
    return rows


def read_workbook(path: Path) -> WorkbookData:
    source = path.resolve()
    if source.suffix.lower() != ".xlsx":
        raise WorkbookContractError("只支持 .xlsx 文件。")
    if not source.is_file():
        raise WorkbookContractError(f"找不到 Excel 文件：{source}")

    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        if workbook.sheetnames != SHEET_NAMES:
            raise WorkbookContractError(
                f"工作表不匹配。预期 {SHEET_NAMES}，实际 {workbook.sheetnames}。"
            )

        tables = {
            name: _read_rows(workbook[name], headers)
            for name, headers in TABLE_COLUMNS.items()
        }
        products = [row for row in tables["产品主表"] if row.get("导入状态") != "示例不导入"]

        return WorkbookData(
            source=source,
            sheet_names=list(workbook.sheetnames),
            products=products,
            parameters=tables["产品参数"],
            images=tables["图片清单"],
            claims=tables["宣称与证据"],
        )
    finally:
        workbook.close()
